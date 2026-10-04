# Builds DSA_Tracker.xlsx: the template for the "DSA Tracker" Google Sheet.
# Run: uv run --with openpyxl python tools/build_sheet.py, then upload to Drive (converts to a Sheet).
# Problems!A1 pulls progress/tracker.csv with IMPORTDATA, so the sheet updates itself.
# Use whole-column ranges (K:K), because open-ended ones (K2:K) break in the xlsx-to-Sheets conversion.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule

URL = "https://raw.githubusercontent.com/Harsh251005/data-structures-and-algorithms/main/progress/tracker.csv"
FONT = "Arial"
HEAD_BG, HEAD_FG = "1F2937", "FFFFFF"
ACCENT = "C2410C"  # saffron, matching DSA Forge
thin = Side(style="thin", color="D1D5DB")

wb = Workbook()
# ---------- Problems ----------
ws = wb.active
ws.title = "Problems"
ws["A1"] = f'=IMPORTDATA("{URL}")'
widths = {"A": 5, "B": 12, "C": 7, "D": 30, "E": 11, "F": 28, "G": 34, "H": 9, "I": 7, "J": 11,
          "K": 6, "L": 14, "M": 14, "N": 10, "O": 60, "P": 44, "Q": 44}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
for col in widths:
    cell = ws[f"{col}1"]
    cell.font = Font(name=FONT, bold=True, color=HEAD_FG)
    cell.fill = PatternFill("solid", fgColor=HEAD_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 30
for col in "DFGO":
    ws.column_dimensions[col].alignment = Alignment(vertical="top", wrap_text=True)
ws.freeze_panes = "E2"
ws.auto_filter.ref = "A1:Q1000"

def fill(hex_):
    return PatternFill("solid", fgColor=hex_, bgColor=hex_)
cf = ws.conditional_formatting
cf.add("E2:E1000", FormulaRule(formula=['$E2="Easy"'], fill=fill("D1FAE5"), font=Font(color="065F46", bold=True)))
cf.add("E2:E1000", FormulaRule(formula=['$E2="Medium"'], fill=fill("FEF3C7"), font=Font(color="92400E", bold=True)))
cf.add("E2:E1000", FormulaRule(formula=['$E2="Hard"'], fill=fill("FEE2E2"), font=Font(color="991B1B", bold=True)))
cf.add("J2:J1000", FormulaRule(formula=['$J2="Yes"'], font=Font(color="065F46", bold=True)))
cf.add("N2:N1000", FormulaRule(formula=['$N2="Yes"'], fill=fill("D1FAE5"), font=Font(color="065F46", bold=True)))
# a re-solve that is due or overdue
cf.add("M2:M1000", FormulaRule(formula=['AND($M2<>"",$N2="No",$M2<=TODAY())'], fill=fill("FFEDD5"), font=Font(color=ACCENT, bold=True)))
cf.add("A2:Q1000", FormulaRule(formula=['AND($A2<>"",MOD(ROW(),2)=0)'], fill=fill("F9FAFB")))

# ---------- Summary ----------
sm = wb.create_sheet("Summary")
P = "Problems"
sm.column_dimensions["A"].width = 30
for c in "BCDE":
    sm.column_dimensions[c].width = 16
sm["A1"] = "DSA progress"
sm["A1"].font = Font(name=FONT, size=16, bold=True, color=ACCENT)
sm["A2"] = "Live from the GitHub repo. The Problems tab refreshes on its own."
sm["A2"].font = Font(name=FONT, size=9, italic=True, color="6B7280")

overall = [
    ("Problems solved", f"=MAX(0,COUNTA({P}!A:A)-1)", "0"),
    ("Total XP", f"=SUM({P}!K:K)", "#,##0"),
    ("Total practice time (hours)", f"=SUM({P}!H:H)/60", "0.0"),
    ("Average minutes per problem", f"=IF(B4=0,0,SUM({P}!H:H)/B4)", "0"),
    ("Solved without hints", f"=IF(B4=0,0,COUNTIF({P}!J:J,\"Yes\")/B4)", "0%"),
    ("Mastered (3 of 3 re-solves)", f"=COUNTIF({P}!N:N,\"Yes\")", "0"),
    ("Re-solves due today or overdue", f"=COUNTIFS({P}!M:M,\"<=\"&TODAY(),{P}!N:N,\"No\")", "0"),
]
sm["A3"], sm["B3"] = "Overall", "Value"
for i, (label, formula, fmt) in enumerate(overall, start=4):
    sm[f"A{i}"] = label
    sm[f"B{i}"] = formula
    sm[f"B{i}"].number_format = fmt

r0 = 4 + len(overall) + 1  # blank row, then the difficulty table
sm[f"A{r0}"], sm[f"B{r0}"], sm[f"C{r0}"], sm[f"D{r0}"], sm[f"E{r0}"] = "By difficulty", "Solved", "Avg minutes", "No-hint rate", "Share"
for j, d in enumerate(["Easy", "Medium", "Hard"], start=r0 + 1):
    sm[f"A{j}"] = d
    sm[f"B{j}"] = f"=COUNTIF({P}!$E:$E,$A{j})"
    sm[f"C{j}"] = f"=IF(B{j}=0,0,AVERAGEIF({P}!$E:$E,$A{j},{P}!$H:$H))"
    sm[f"D{j}"] = f"=IF(B{j}=0,0,COUNTIFS({P}!$E:$E,$A{j},{P}!$J:$J,\"Yes\")/B{j})"
    sm[f"E{j}"] = f"=IF($B$4=0,0,B{j}/$B$4)"
    sm[f"C{j}"].number_format = "0"
    sm[f"D{j}"].number_format = "0%"
    sm[f"E{j}"].number_format = "0%"

for hdr_row, last_col in ((3, "B"), (r0, "E")):
    for col in "ABCDE"[: "ABCDE".index(last_col) + 1]:
        cell = sm[f"{col}{hdr_row}"]
        cell.font = Font(name=FONT, bold=True, color=HEAD_FG)
        cell.fill = PatternFill("solid", fgColor=HEAD_BG)
        cell.alignment = Alignment(horizontal="left" if col == "A" else "center")
for row in sm.iter_rows(min_row=4, max_row=r0 + 3, max_col=5):
    for cell in row:
        if cell.row != r0 and cell.value is not None:
            cell.font = Font(name=FONT, size=11)
            cell.border = Border(bottom=thin)
            if cell.column > 1:
                cell.alignment = Alignment(horizontal="center")

wb.move_sheet("Summary", offset=-1)  # Summary first
wb.active = 0
wb.save("DSA_Tracker.xlsx")  # written to the current directory
print("ok")
